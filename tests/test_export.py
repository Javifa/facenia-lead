"""
Pruebas para los servicios de exportación a Excel y CSV.
"""

import io
import openpyxl
from export_service import export_leads_to_excel_bytes, export_leads_to_csv_bytes
from config import WEB_STATUS_NO_WEB, WEB_STATUS_FOUND, PRIORITY_HIGH, PRIORITY_LOW


def test_excel_export_structure():
    leads = [
        {
            "name": "Reformas Alfa",
            "sector": "Reformas y construcción",
            "web_status": WEB_STATUS_NO_WEB,
            "priority": PRIORITY_HIGH,
            "score": 90,
            "commercial_status": "Pendiente de contactar"
        },
        {
            "name": "Clínica Beta",
            "sector": "Clínicas y centros sanitarios",
            "web_status": WEB_STATUS_FOUND,
            "priority": PRIORITY_LOW,
            "score": 40,
            "commercial_status": "Pendiente de revisar"
        }
    ]

    excel_bytes = export_leads_to_excel_bytes(leads)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames

    assert "Todos los Prospectos" in sheet_names
    assert "Sin Web Indicada" in sheet_names
    assert "Prioridad Alta" in sheet_names
    assert "Pendientes de Verificar" in sheet_names

    ws_no_web = wb["Sin Web Indicada"]
    assert ws_no_web.max_row == 2  # Fila 1 = cabecera, Fila 2 = 1 lead


def test_csv_export_notion_compatibility():
    leads = [
        {
            "name": "Instalaciones Málaga & Compañía",
            "sector": "Instalaciones y mantenimiento",
            "phone": "+34 952 11 22 33",
            "web_status": WEB_STATUS_NO_WEB,
            "priority": PRIORITY_HIGH,
            "score": 85,
            "commercial_status": "Pendiente de contactar"
        }
    ]

    csv_bytes = export_leads_to_csv_bytes(leads)
    assert len(csv_bytes) > 0

    # Debe comenzar con el BOM UTF-8 (b'\xef\xbb\xbf') para compatibilidad nativa en Windows y Notion
    assert csv_bytes.startswith(b"\xef\xbb\xbf")

    decoded_text = csv_bytes.decode("utf-8-sig")
    assert "Nombre Comercial" in decoded_text
    assert "Instalaciones Málaga & Compañía" in decoded_text
    assert "Estado Comercial" in decoded_text
